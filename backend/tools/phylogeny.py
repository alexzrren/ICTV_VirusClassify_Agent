"""
Phylogenetic placement tool via EPA-ng.

Wraps the full EPA-ng pipeline (HMM marker extraction → MAFFT alignment →
raxml-ng model evaluation → EPA-ng placement → LCA classification) into a
single agent-callable function.

Reference trees and MSAs come from the pre-built ictv_classifier project.
RNA viruses use RdRp markers; DNA viruses use Rep/Pol/E1 markers.

Requires in PATH: getorf, hmmsearch, mafft, raxml-ng, epa-ng.
Requires Python: ete4, openpyxl, Biopython.
"""

from __future__ import annotations

import json
import os
import re
import shutil
import statistics
import subprocess
import tempfile
import sqlite3
from pathlib import Path
from typing import Optional

import openpyxl

# ── Paths ────────────────────────────────────────────────────────────────────

_REF_ROOT = Path("/home/renzirui/Projects/ICTV/ictv_classifier/reference")
_HMM_RNA_DIR = Path("/home/renzirui/Projects/Phylogenetic_Background/VHDB_RdRp_firstpass/Crop_Profile")
_HMM_DNA_DIR = Path("/home/renzirui/Projects/Phylogenetic_Background/Rep_Marker_Profile/Crop_Profile")
_CROP_BY_HMM = "/home/renzirui/Projects/ChinaBatVirome/ViralFamilyPhylo/crop_by_hmm.py"
_TAX_DB = Path(__file__).resolve().parent.parent.parent / "data" / "taxonomy.db"
_MICROMAMBA_BIN = str(Path.home() / "micromamba" / "bin")

def _env_with_path():
    """Return an environment dict with micromamba/bin prepended to PATH."""
    env = os.environ.copy()
    existing = env.get("PATH", "")
    if _MICROMAMBA_BIN not in existing:
        env["PATH"] = f"{_MICROMAMBA_BIN}:{existing}"
    return env

# Which families are DNA viruses (use Rep/Pol marker, not RdRp)
DNA_FAMILIES = {
    "Adenoviridae", "Anelloviridae", "Asfarviridae", "Circoviridae",
    "Papillomaviridae", "Parvoviridae", "Polyomaviridae", "Poxviridae",
}


def _resolve_paths(family: str) -> tuple[Optional[Path], Optional[Path], Optional[Path], str]:
    """Resolve reference tree, MSA, HMM profile for a family.

    Returns (tree_path, msa_path, hmm_path, marker_type)
    or (None, None, None, "") if family not found.
    """
    ref_dir = _REF_ROOT / family
    if not ref_dir.is_dir():
        return None, None, None, ""

    # Tree (IQ-TREE output, either marker_tree or RdRp_tree)
    tree = None
    for cand in ["marker_tree.treefile", "RdRp_tree.treefile"]:
        p = ref_dir / cand
        if p.exists() and p.stat().st_size > 100:
            tree = p; break
    if not tree:
        return None, None, None, ""

    # MSA (marker.afa or RdRp.afa)
    msa = None
    for cand in ["marker.afa", "RdRp.afa"]:
        p = ref_dir / cand
        if p.exists() and p.stat().st_size > 100:
            msa = p; break
    if not msa:
        return None, None, None, ""

    # HMM profile
    marker_type = "Rep" if family in DNA_FAMILIES else "RdRp"
    hmm_dir = _HMM_DNA_DIR if family in DNA_FAMILIES else _HMM_RNA_DIR
    hmm = hmm_dir / f"{family}.hmm"
    if not hmm.exists():
        return tree, msa, None, marker_type

    return tree, msa, hmm, marker_type


# ── Pipeline steps ───────────────────────────────────────────────────────────

def _extract_marker(query_nt: str, hmm_path: Path, marker_type: str) -> Optional[str]:
    """Extract marker protein from nucleotide query via HMM.

    Returns the protein sequence (FASTA string) or None.
    """
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        query_fa = tmp / "query.fna"
        query_fa.write_text(f">query\n{query_nt}\n")

        out_faa = tmp / "marker.faa"
        dom_csv = tmp / "domtbl.csv"
        crop_csv = tmp / "crop.csv"

        env = _env_with_path()
        try:
            subprocess.run(
                ["python3", _CROP_BY_HMM,
                 "-i", str(query_fa), "-m", str(hmm_path),
                 "-o", str(out_faa), "-d", str(dom_csv), "-r", str(crop_csv)],
                capture_output=True, text=True, timeout=120, env=env,
            )
        except subprocess.TimeoutExpired:
            return None

        if not out_faa.exists() or out_faa.stat().st_size < 10:
            return None

        marker_text = out_faa.read_text().strip()
        lines = marker_text.split("\n")
        seq = "".join(l.strip() for l in lines if not l.startswith(">"))
        if len(seq) < 50:
            return None
        return marker_text  # full FASTA


def _clean_fasta_headers(fasta_text: str) -> str:
    """Strip everything after the first token in each FASTA header."""
    cleaned = []
    for line in fasta_text.split("\n"):
        if line.startswith(">"):
            cleaned.append(">" + line[1:].split()[0])
        else:
            cleaned.append(line.strip())
    return "\n".join(cleaned)


def _align_query(marker_faa: str, ref_msa_path: Path) -> Optional[str]:
    """Align query marker to reference MSA with MAFFT --add --keeplength.

    Returns the aligned query sequence string (no header, gaps included).
    """
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        query_fa = tmp / "query.faa"
        query_fa.write_text(_clean_fasta_headers(marker_faa))

        result = subprocess.run(
            ["mafft", "--add", str(query_fa), "--keeplength", str(ref_msa_path)],
            capture_output=True, text=True, timeout=120, env=_env_with_path(),
        )
        if result.returncode != 0:
            return None

        # MAFFT writes combined alignment to stdout
        msa_text = result.stdout
        if not msa_text:
            return None

        # Parse interleaved FASTA
        seqs = {}
        cur = None
        for line in msa_text.split("\n"):
            if line.startswith(">"):
                cur = line[1:].split()[0]
                seqs[cur] = ""
            elif cur:
                seqs[cur] += line.strip()

        # Return the query (from cleaned header, first token of first seq)
        for hdr, seq in seqs.items():
            marker_hdr = _clean_fasta_headers(marker_faa).split("\n")[0][1:].split()[0]
            if marker_hdr == hdr or hdr == "query":
                return seq
    return None


def _run_raxml(ref_msa_path: Path, ref_tree_path: Path, workdir: Path) -> Optional[Path]:
    """Run raxml-ng --evaluate on reference MSA.

    Returns path to bestModel file or None.
    """
    # Clean reference MSA headers for compatibility
    cleaned_msa = workdir / "ref_clean.afa"
    with open(ref_msa_path) as fin, open(cleaned_msa, "w") as fout:
        for line in fin:
            if line.startswith(">"):
                fout.write(">" + line[1:].split()[0] + "\n")
            else:
                fout.write(line.strip() + "\n")

    env = _env_with_path()
    env.setdefault("OMP_NUM_THREADS", "2")
    result = subprocess.run(
        ["raxml-ng", "--evaluate",
         "--msa", str(cleaned_msa),
         "--model", "LG+I+G4",
         "--tree", str(ref_tree_path),
         "--brlen", "scaled",
         "--threads", "2"],
        capture_output=True, text=True, timeout=300, env=env, cwd=str(workdir),
    )
    model_file = workdir / "ref_clean.afa.raxml.bestModel"
    if model_file.exists():
        return model_file
    return None


def _run_epa(msa_path: Path, tree_path: Path, model_path: Path,
             query_aligned: str, query_name: str, workdir: Path) -> Optional[Path]:
    """Run EPA-ng placement. Returns path to epa_result.jplace."""
    # Write aligned query
    query_fa = workdir / "query_aligned.faa"
    query_fa.write_text(f">{query_name}\n{query_aligned}\n")

    epa_out = workdir / "epa_out"
    epa_out.mkdir(exist_ok=True)

    subprocess.run(
        ["epa-ng", "--ref-msa", str(msa_path), "--tree", str(tree_path),
         "--model", str(model_path), "--query", str(query_fa),
         "--outdir", str(epa_out)],
        capture_output=True, text=True, timeout=120, env=_env_with_path(),
    )

    jplace = epa_out / "epa_result.jplace"
    return jplace if jplace.exists() else None


# ── LCA Classification ──────────────────────────────────────────────────────

TAX_LEVELS = ["Realm", "Kingdom", "Phylum", "Class", "Order",
              "Family", "Subfamily", "Genus", "Subgenus", "Species"]


def _load_vmr_flat(xlsx_path: Path) -> dict[str, dict]:
    """Build bare-accession -> taxonomy mapping from VMR Excel."""
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True, data_only=True)
    ws = wb.active
    headers = [str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    acc_col = None; tax_cols = {}
    for i, h in enumerate(headers):
        if h and "genbank" in h.lower(): acc_col = i
        if h in TAX_LEVELS: tax_cols[h] = i

    if acc_col is None:
        # Fallback: known VMR column positions
        acc_col = 23
        tax_cols = {lv: i for i, lv in enumerate(TAX_LEVELS) if i < 28}
        for i, h in enumerate(headers):
            if h in TAX_LEVELS: tax_cols[h] = i

    acc2tax = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        raw = row[acc_col] if acc_col < len(row) else ""
        if not raw: continue
        taxonomy = {lv: (str(row[tax_cols[lv]]).strip() if lv in tax_cols and tax_cols[lv] < len(row) and row[tax_cols[lv]] else "")
                    for lv in TAX_LEVELS}
        for part in str(raw).split(";"):
            part = part.strip()
            m = re.search(r'([A-Z]{1,2}_?\d{5,9})', part)
            if m:
                acc2tax[m.group(1)] = taxonomy
    return acc2tax


def _map_tips_to_leaves(tree):
    """Return {leaf_name_clean: ete4_node}."""
    return {n.name: n for n in tree.leaves()}


def _compute_lca(tax_list):
    """Deepest level where all taxa agree."""
    if not tax_list:
        return {l: "" for l in TAX_LEVELS}
    result = {}
    conflict = False
    for lv in TAX_LEVELS:
        if conflict:
            result[lv] = ""; continue
        vals = {t.get(lv, "") for t in tax_list if t.get(lv, "")}
        if len(vals) == 1:
            result[lv] = vals.pop()
        else:
            result[lv] = ""; conflict = True
    return result


def phylogenetic_placement(query_nt: str, family: str) -> dict:
    """Run full EPA-ng phylogenetic placement pipeline.

    Args:
        query_nt: Query nucleotide genome sequence.
        family: Virus family name (e.g. "Papillomaviridae").

    Returns:
        Structured classification dict or {"error": "..."}.
    """
    tree_path, msa_path, hmm_path, marker_type = _resolve_paths(family)
    if not tree_path or not msa_path:
        return {"error": f"No reference tree/MSA for {family}. "
                         f"Available families: {sorted(d.name for d in _REF_ROOT.iterdir() if d.is_dir())}"}
    if not hmm_path:
        return {"error": f"No HMM profile for {family} (marker type: {marker_type})"}

    vmr_xlsx = tree_path.parent / f"ICTV_{family}_VMR.xlsx"
    if not vmr_xlsx.exists():
        # Try glob
        candidates = list(tree_path.parent.glob("ICTV_*_VMR.xlsx"))
        if candidates:
            vmr_xlsx = candidates[0]
        else:
            return {"error": f"No VMR xlsx found in {tree_path.parent}"}

    workdir = Path(tempfile.mkdtemp(prefix=f"epa_{family}_"))

    try:
        # Step 1: Extract marker
        marker_faa = _extract_marker(query_nt, hmm_path, marker_type)
        if not marker_faa:
            return {"error": f"Failed to extract {marker_type} marker from query sequence "
                             f"(genome may be too short or missing the marker gene)"}

        marker_aa = "".join(l.strip() for l in marker_faa.split("\n") if not l.startswith(">"))

        # Step 2: Align query to reference MSA
        query_aligned = _align_query(marker_faa, msa_path)
        if not query_aligned:
            return {"error": "MAFFT alignment failed"}

        # Step 3: raxml-ng evaluate
        model_path = _run_raxml(msa_path, tree_path, workdir)
        if not model_path:
            return {"error": "raxml-ng model evaluation failed"}

        # Step 4: EPA-ng placement
        jplace = _run_epa(workdir / "ref_clean.afa", tree_path, model_path,
                          query_aligned, "query", workdir)
        if not jplace:
            return {"error": "EPA-ng placement failed"}

        # Step 5: Parse placement + LCA classify
        import json as _json
        from ete4 import Tree

        with open(jplace) as f:
            jp = _json.load(f)

        tree_str = jp["tree"]
        edge_nums = [int(m.group(1)) for m in re.finditer(r'\{(\d+)\}', tree_str)]
        clean_tree = re.sub(r'\{(\d+)\}', '', tree_str)
        tree = Tree(clean_tree, parser=1)

        # Edge number → node mapping (postorder, exclude root)
        nodes_po = [n for n in tree.traverse("postorder") if n.up is not None]
        edge2node = dict(zip(edge_nums, nodes_po))

        # Placement
        p_info = jp["placements"][0]
        best = p_info["p"][0]  # [edge_num, likelihood, LWR, distal, pendant]
        edge_num = int(best[0])
        lwr = float(best[2])
        pendant = float(best[4]) if len(best) > 4 else 0.0

        # Taxonomy for tips below placement
        acc2tax = _load_vmr_flat(vmr_xlsx)
        node = edge2node.get(edge_num)
        taxa_below = []
        if node:
            for leaf in node.leaves():
                bare = leaf.name.split(".")[0] if "." in leaf.name else leaf.name
                tax = acc2tax.get(bare)
                if tax:
                    taxa_below.append(tax)

        lca = _compute_lca(taxa_below) if taxa_below else {l: "" for l in TAX_LEVELS}

        # Novel species assessment
        ref_pendants = [n.dist for n in tree.leaves()]
        ref_mean = statistics.mean(ref_pendants) if ref_pendants else 0.0
        ref_std = statistics.stdev(ref_pendants) if len(ref_pendants) > 1 else 0.0
        pendant_thr = ref_mean + 2 * ref_std

        novel_reasons = []
        if lca.get("Genus") and not lca.get("Species"):
            novel_reasons.append("species_conflict")
        if pendant > pendant_thr:
            novel_reasons.append(f"long_pendant ({pendant:.4f} > {pendant_thr:.4f})")

        is_novel = bool(novel_reasons)
        if lwr >= 0.8:
            conf = "High"
        elif lwr >= 0.5:
            conf = "Medium"
        else:
            conf = "Low"

        return {
            "family": family,
            "marker": marker_type,
            "marker_length_aa": len(marker_aa),
            "ref_tips": len(list(tree.leaves())),
            "best_placement": {
                "edge_num": edge_num,
                "LWR": round(lwr, 4),
                "pendant_length": round(pendant, 4),
                "pendant_threshold": round(pendant_thr, 4),
            },
            "lca_classification": lca,
            "novel_species": is_novel,
            "novel_reasons": novel_reasons,
            "confidence": conf,
            "note": (
                f"Phylogenetic placement on {family} reference tree ({marker_type} marker, "
                f"{len(list(tree.leaves()))} tips). Closest placement edge #{edge_num} "
                f"with LWR={lwr:.3f}."
            ),
        }

    finally:
        shutil.rmtree(workdir, ignore_errors=True)


def list_available_placement_families() -> dict[str, dict]:
    """Return dict of families with usable reference trees."""
    result = {}
    for d in sorted(_REF_ROOT.iterdir()):
        if not d.is_dir(): continue
        fam = d.name
        tree, msa, _, _ = _resolve_paths(fam)
        if tree and msa:
            tips = len(re.findall(r'[A-Z]{1,2}[0-9]{5,9}', tree.read_text()))
            result[fam] = {
                "tips": tips,
                "marker": "Rep" if fam in DNA_FAMILIES else "RdRp",
                "tree_exists": True,
            }
    return result
